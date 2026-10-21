"""Build the Vietnamese Word report, organized by exercise and test evidence."""
import json
import shutil
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
report = Document()
section = report.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2), Cm(2)
section.left_margin, section.right_margin = Cm(2.5), Cm(2)
normal = report.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)
for name in ("Title", "Heading 1", "Heading 2"):
    report.styles[name].font.name = "Times New Roman"
    report.styles[name].font.color.rgb = RGBColor.from_string("17365D")
section.header.paragraphs[0].text = "LAB 5 — LOCAL MQTT GATEWAY"
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Trang ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def paragraph(text, style=None):
    return report.add_paragraph(text, style)


def code(text):
    p = paragraph("")
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_together = False
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "F1F5F9")
    p._p.get_or_add_pPr().append(shade)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def table(headers, rows):
    t = report.add_table(rows=1, cols=len(headers))
    t.style = "Light Shading Accent 1"
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)


def excerpt(path, limit=45):
    lines = (ROOT / path).read_text(encoding="utf-8-sig").splitlines()
    if len(lines) > limit:
        lines = lines[:limit] + ["# ... Xem tệp nguồn để đọc đầy đủ."]
    paragraph("Tệp: " + path)
    code("\n".join(lines))


manifest = json.loads((ROOT / "scripts/exercises.json").read_text(encoding="utf-8"))
results = json.loads((ROOT / "reports/test-results.json").read_text(encoding="utf-8"))
ui_path = ROOT / "reports/dashboard-results.json"
ui = json.loads(ui_path.read_text(encoding="utf-8")) if ui_path.exists() else None
tested = datetime.fromisoformat(results["tested_at"]).astimezone(timezone(timedelta(hours=7)))
report.core_properties.title = "Báo cáo thực hành Lab 5 — Local MQTT Gateway"
report.core_properties.subject = "5 bài trên lớp và 3 bài về nhà"

paragraph("BÁO CÁO THỰC HÀNH\nLAB 5", "Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph("LOCAL MQTT BROKER VÀ IOT GATEWAY", "Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph("\n")
for label in ("Họ và tên: ............................................................",
              "Mã số sinh viên: .....................................................",
              "Lớp: .......................................................................",
              "Giảng viên: .............................................................."):
    paragraph(label)
paragraph("\nNội dung: 5 bài thực hành trên lớp và 3 bài thực hành về nhà.")
paragraph("Đề bài: Lab 5_1.docx.")
paragraph("Kiểm thử MQTT: " + tested.strftime("%d/%m/%Y %H:%M:%S (UTC+7)"))
paragraph("Báo cáo sử dụng bằng chứng kiểm thử đã ghi nhận. CPU/RAM là số đo thật; cảm biến là mô phỏng.")
report.add_page_break()

report.add_heading("I. Môi trường, cấu trúc và kiến trúc", 1)
paragraph("Mục tiêu: xây dựng gateway MQTT cục bộ gồm broker, client CLI/Python, gửi số đo hệ thống, "
          "nhận lệnh, tổng hợp dữ liệu cảm biến, xác thực/phân quyền và dashboard Node-RED.")
table(["Thành phần", "Môi trường thực nghiệm"], [
    ["Hệ điều hành", "Windows x64; C:\\IOT\\lab5"],
    ["Python và thư viện", "Python 3.10.11; .venv; paho-mqtt 2.1.0; psutil 7.2.2"],
    ["Broker", "Mosquitto 2.1.2; 127.0.0.1:1883"],
    ["Dashboard", "Node.js 24.15.0; Node-RED 5.0.8; FlowFuse Dashboard 1.33.0"],
])
paragraph("Chưa thực nghiệm trên Raspberry Pi hoặc cảm biến/thiết bị chấp hành thật. "
          "Hướng dẫn triển khai Linux tương ứng nằm trong README.md.")
report.add_heading("1. Cấu trúc thư mục", 2)
table(["Thư mục", "Nội dung"], [[e["folder"], e["title"]] for e in manifest])
paragraph("Mỗi bài có README.md và mã nguồn/cấu hình/lệnh chạy tương ứng. gateway/common.py, .venv, "
          "config, tools và môi trường node-red được dùng chung. Các module gateway cũ vẫn gọi được mã nguồn "
          "trong thư mục từng bài. Báo cáo chính ở Bao_cao/Bao_cao_Lab_5.docx.")
report.add_heading("2. Luồng dữ liệu", 2)
table(["Nguồn → đích", "Topic / chu kỳ"], [
    ["Sensor → Mosquitto → Aggregator", "sensor/+/data"],
    ["CPU/RAM → Mosquitto → Node-RED", "gateway/system/metrics — 5 giây"],
    ["Aggregator → Mosquitto → Node-RED", "gateway/aggregated_report — 10 giây"],
    ["Operator → Mosquitto → Subscriber", "gateway/command/#"],
    ["CLI publisher → CLI subscriber", "gateway/test"],
])

snippets = {
    "Tren_lop/Bai_1": ["mosquitto.conf"], "Tren_lop/Bai_2": ["nhan.ps1", "gui.ps1"],
    "Tren_lop/Bai_3": ["check_connection.py"], "Tren_lop/Bai_4": ["metrics_publisher.py"],
    "Tren_lop/Bai_5": ["command_subscriber.py"], "Ve_nha/Bai_1": ["aggregator.py"],
    "Ve_nha/Bai_2": ["mosquitto-secure.conf.example", "acl"],
}
for index, item in enumerate(manifest, 1):
    report.add_page_break()
    report.add_heading(f"II.{index}. {item['title']}", 1)
    report.add_heading("1. Mục tiêu", 2)
    paragraph(item["objective"])
    report.add_heading("2. Tệp thực hiện", 2)
    paragraph("Thư mục: " + item["folder"])
    for filename in item["files"]:
        paragraph(filename, "List Bullet")
    report.add_heading("3. Cách thực hiện và chạy", 2)
    paragraph("Chạy các lệnh từ C:\\IOT\\lab5. Môi trường chung đã cài; có thể cài lại bằng scripts/setup.ps1.")
    code("\n\n".join(item["commands"]))
    paragraph(item["notes"])
    report.add_heading("4. Mã nguồn/cấu hình tiêu biểu", 2)
    for filename in snippets.get(item["folder"], []):
        excerpt(item["folder"] + "/" + filename)
    if item["folder"] == "Tren_lop/Bai_3":
        paragraph("Session ở gateway/common.py thiết lập callback API v2, xác thực và vòng lặp mạng. "
                  "Chương trình chờ CONNACK thành công; khi đăng ký topic, chờ thêm SUBACK.")
    if item["folder"] == "Ve_nha/Bai_1":
        paragraph("Vòng lặp main chờ bằng đồng hồ monotonic, gọi snapshot mỗi 10 giây và publish bằng QoS 1. "
                  "Nếu gửi lỗi, restore đưa dữ liệu về bộ đệm. node_count đếm node; message_count đếm bản tin hợp lệ "
                  "đã nhận, gồm cả cập nhật nhiều lần của cùng node.")
    if item["folder"] == "Ve_nha/Bai_2":
        paragraph("cau_hinh_bao_mat.py tạo mật khẩu ngẫu nhiên hoặc giữ tài khoản đã có; "
                  "mosquitto_passwd -U băm mật khẩu bằng công cụ chính thức. Tài khoản thực không được nhúng vào báo cáo.")
    if item["folder"] == "Ve_nha/Bai_3":
        paragraph("MQTT in gateway/system/metrics → function tách CPU/RAM → 2 ui-gauge và ui-chart. "
                  "MQTT in gateway/aggregated_report → function tóm tắt → ui-text và ui-table. "
                  "Flow không chứa mật khẩu; thông tin xác thực broker được lưu riêng trong môi trường node-red.")
    report.add_heading("5. Kết quả và bằng chứng", 2)
    paragraph("Kết quả cần quan sát: " + item["expected"])
    for result in results["results"]:
        if not any(result["name"].startswith(p) for p in item["test_prefixes"]):
            continue
        paragraph(("ĐẠT — " if result["passed"] else "CHƯA ĐẠT — ") + result["name"])
        evidence = result["evidence"]
        code(evidence if isinstance(evidence, str) else json.dumps(evidence, indent=2, ensure_ascii=False))
    if item["folder"] == "Ve_nha/Bai_3" and ui:
        paragraph(f"Kiểm thử trình duyệt: HTTP {ui['httpStatus']}; {len(ui['errors'])} lỗi JavaScript; "
                  "giao diện có Gauge CPU/RAM, biểu đồ và dữ liệu node1/node2.")
        paragraph("Gauge tại thời điểm chụp: " + ", ".join(ui.get("gaugeValues", [])))
        if (ROOT / "reports/dashboard.png").exists():
            report.add_picture(str(ROOT / "reports/dashboard.png"), width=Cm(16))
            paragraph("Hình 1. Dashboard với CPU/RAM thật và hai cảm biến mô phỏng.", "Caption")

report.add_page_break()
report.add_heading("III. Tổng hợp kiểm thử và cách chạy lại", 1)
table(["Kiểm thử", "Kết quả"], [[r["name"], "ĐẠT" if r["passed"] else "CHƯA ĐẠT"] for r in results["results"]])
paragraph(f"Số kiểm thử MQTT đạt: {sum(r['passed'] for r in results['results'])}/{len(results['results'])}. "
          "Broker kiểm thử dùng cổng riêng 18884/18885. Bằng chứng gốc: reports/test-results.json, "
          "reports/dashboard-results.json và reports/dashboard.png.")
if results["error"]:
    paragraph("Lỗi trong lần kiểm thử: " + results["error"])
code(".\\.venv\\Scripts\\python.exe tests\\integration.py\nnode tests\\dashboard.cjs\n"
     ".\\.venv\\Scripts\\python.exe scripts\\build_report.py")
paragraph("Kiểm thử dashboard cần hệ thống demo đang chạy, playwright-core và trình duyệt Chromium tương thích.")
code("# Xem trạng thái\n.\\.venv\\Scripts\\python.exe scripts\\manage.py status\n\n"
     "# Dừng / khởi động\npowershell -ExecutionPolicy Bypass -File scripts\\stop.ps1\n"
     "powershell -ExecutionPolicy Bypass -File scripts\\start.ps1")
report.add_heading("IV. Đánh giá và giới hạn", 1)
paragraph("Các bài đã triển khai broker cục bộ, truyền nhận CLI/Python, đo CPU/RAM theo chu kỳ, nhận lệnh, "
          "tổng hợp nhiều cảm biến, xác thực/ACL và hiển thị dashboard Node-RED.")
paragraph("Chưa kiểm thử trên Linux/Raspberry Pi hoặc giữa các thiết bị LAN. MQTT 1883 chưa dùng TLS. "
          "Bộ đệm ở RAM không lưu bền qua mất điện. QoS 1 có thể giao lại bản tin; chưa khử trùng lặp theo ID ứng dụng. "
          "Các sensor là mô phỏng, chưa bật/tắt LED vật lý.")
paragraph("Kết quả npm audit đã lưu ghi nhận 5 vấn đề trong các thư viện bundle của gói npm phụ thuộc Node-RED; "
          "npm audit fix chưa tự sửa được. Xem reports/npm-audit.json. Editor/dashboard chỉ lắng nghe 127.0.0.1.")
report.add_heading("V. Tài liệu tham khảo", 1)
for url in ("https://mosquitto.org/download/", "https://mosquitto.org/documentation/authentication-methods/",
            "https://mosquitto.org/man/mosquitto-conf-5.html", "https://eclipse.dev/paho/clients/python/docs/",
            "https://nodered.org/docs/getting-started/windows", "https://dashboard.flowfuse.com/getting-started.html"):
    paragraph(url)

target = ROOT / "Bao_cao/Bao_cao_Lab_5.docx"
target.parent.mkdir(exist_ok=True)
report.save(target)
shutil.copyfile(target, ROOT / "reports/Bao_cao_Lab_5.docx")
print(f"Report saved: {target}")
